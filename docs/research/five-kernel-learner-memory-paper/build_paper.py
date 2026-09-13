"""Build an editable Chinese academic manuscript with the bundled Python runtime.

The Markdown is the content authority; render with the bundled documents renderer.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE=Path(__file__).resolve().parent


def font(run, size=None, bold=None):
    run.font.name='Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)


def hyperlink(p,text,url):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'),'Times New Roman'); fonts.set(qn('w:eastAsia'),'Songti SC')
    props.append(fonts)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'19');props.append(size)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'000000'); props.append(color)
    run.append(props); t=OxmlElement('w:t');t.text=text;run.append(t);h.append(run);p._p.append(h)


def inline(p,text):
    pattern=r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))'
    for chunk in re.split(pattern,text):
        if not chunk: continue
        if chunk.startswith('**'):
            r=p.add_run(chunk[2:-2]);font(r,bold=True)
        elif chunk.startswith('`'):
            r=p.add_run(chunk[1:-1]);font(r)
        elif re.fullmatch(r'\[[^\]]+\]\([^\)]+\)',chunk):
            label,url=re.match(r'\[([^\]]+)\]\(([^\)]+)\)',chunk).groups()
            hyperlink(p,label,url)
        else: font(p.add_run(chunk))


def paragraph(doc,text,style=None):
    p=doc.add_paragraph(style=style); inline(p,text); return p


def add_table(doc,rows):
    cols=len(rows[0]);t=doc.add_table(rows=1,cols=cols)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={4:[2.45,4.35,5.15,5.24],6:[1.35,3.45,2.9,2.9,3.1,3.49]}.get(cols,[17.19/cols]*cols)
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for ridx,row in enumerate(rows):
        cells=t.rows[0].cells if ridx==0 else t.add_row().cells
        trpr=t.rows[ridx]._tr.get_or_add_trPr()
        split=OxmlElement('w:cantSplit');trpr.append(split)
        if ridx==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for cidx,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[cidx]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for name,val in [('top','75'),('bottom','75'),('left','75'),('right','75')]:
                el=OxmlElement('w:'+name);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            if ridx==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F2F3');pr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.1
            p.paragraph_format.first_line_indent=Cm(0)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if cols==6 and cidx!=1 else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,value)
            for r in p.runs: font(r,9.5,ridx==0)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(2)


def build():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    sec.header_distance=Cm(.85);sec.footer_distance=Cm(.85)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Songti SC')
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.widow_control=True
    for s,size in [('Title',19),('Heading 1',14),('Heading 2',12),('Caption',9.5)]:
        st=doc.styles[s];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
        st.font.bold=s!='Caption';st.paragraph_format.space_before=Pt(11 if 'Heading' in s else 0)
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=s!='Caption'
    head=sec.header.paragraphs[0];head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(head.add_run('五核学习者记忆建模与证据治理'),8.5)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    doc.core_properties.title='面向计算机专业教育的五核学习者记忆建模与证据治理'
    doc.core_properties.subject='中文研究论文；系统能力与可审计自动实验'
    doc.core_properties.author=''
    # Bundled default styles contain a decorative Title paragraph border.
    # Remove paragraph borders explicitly; scientific tables keep tcBorders.
    for el in list(doc.styles.element.iter(qn('w:pBdr'))):
        el.getparent().remove(el)
    lines=(HERE/'manuscript_zh.md').read_text().splitlines()
    i=0;refs=False
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s or s.startswith('<!--'): continue
        if s.startswith('|'):
            raw=[s]
            while i<len(lines) and lines[i].strip().startswith('|'):
                raw.append(lines[i].strip());i+=1
            rows=[[x.strip() for x in r.strip('|').split('|')] for r in raw if not re.match(r'^\|\s*:?-',r)]
            add_table(doc,rows);continue
        if s.startswith('!['):
            name,path=re.match(r'!\[([^\]]*)\]\(([^\)]+)\)',s).groups()
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(HERE/path),width=Cm(17.1))
            continue
        if s.startswith('# '):
            p=paragraph(doc,s[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(13);continue
        if s.startswith('## '):
            title=s[3:];refs=title=='参考文献'
            paragraph(doc,title,'Heading 1');continue
        if s.startswith('### '): paragraph(doc,s[4:],'Heading 2');continue
        if re.match(r'^图\d ',s):
            p=paragraph(doc,s,'Caption');p.paragraph_format.keep_with_next=False;continue
        if re.match(r'^表\d ',s):
            p=paragraph(doc,s,'Caption');p.paragraph_format.keep_with_next=True;continue
        p=paragraph(doc,s)
        if refs:
            p.paragraph_format.left_indent=Cm(.65);p.paragraph_format.first_line_indent=Cm(-.65)
            p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
            for r in p.runs:font(r,9.5)
        elif not s.startswith('**关键词'):
            p.paragraph_format.first_line_indent=Pt(22)
    out=HERE/'five_kernel_learner_memory_zh.docx';doc.save(out)
    print(out)


if __name__=='__main__': build()
