"""Build Chinese manuscript and journal-style vector figures from retained data.

Dependencies: python-docx, reportlab; raster previews use sharp (Node).
No new experiments, inference, network access, or synthetic uncertainty.
"""
from pathlib import Path
import argparse, hashlib, json, re, subprocess, os
import xml.etree.ElementTree as ET
from reportlab.graphics.shapes import Drawing, String, Rect, Line, Polygon, Circle
from reportlab.graphics import renderSVG, renderPDF
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

HERE=Path(__file__).resolve().parent
OUT=HERE/'figures_zh'
BLUE='#0072B2'; ORANGE='#E69F00'; STALE='#D55E00'; GRAY='#D9DEE2'; INK='#17191B'; GRID='#DFE3E6'
W=183/25.4*72
FONT_DIR=Path(os.environ.get('PAPER_FONT_DIR','/System/Library/Fonts/Supplemental'))
for name,file in [('Arial Unicode MS','Arial Unicode.ttf'),('Arial','Arial.ttf'),('Arial-Bold','Arial Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(FONT_DIR/file)))

class Figure:
    def __init__(self,h):
        self.h=h;self.d=Drawing(W,h);self.rect(0,0,W,h,'#FFFFFF',None)
    def txt(self,x,y,s,size=7,bold=False,anchor='start',color=INK):
        font='Arial Unicode MS' if re.search(r'[^\x00-\x7f]',str(s)) else ('Arial-Bold' if bold else 'Arial')
        self.d.add(String(x,self.h-y,str(s),fontName=font,fontSize=size,fillColor=HexColor(color),textAnchor=anchor))
    def line(self,x,y,x2,y2,color=INK,width=.5,dash=None):
        self.d.add(Line(x,self.h-y,x2,self.h-y2,strokeColor=HexColor(color),strokeWidth=width,strokeDashArray=dash))
    def rect(self,x,y,w,h,fill=None,stroke=INK,width=.5):
        self.d.add(Rect(x,self.h-y-h,w,h,fillColor=HexColor(fill) if fill else None,strokeColor=HexColor(stroke) if stroke else None,strokeWidth=width))
    def point(self,x,y,color,shape='circle',size=2.4):
        if shape=='square':self.rect(x-size,y-size,2*size,2*size,color,'#FFFFFF',.35)
        else:self.d.add(Circle(x,self.h-y,size,fillColor=HexColor(color),strokeColor=HexColor('#FFFFFF'),strokeWidth=.35))
    def arrow(self,x,y,x2,y2,color=INK):
        self.line(x,y,x2,y2,color,.6)
        if abs(y2-y)<1: pts=[x2,y2,x2-4,y2-2,x2-4,y2+2]
        else:pts=[x2,y2,x2-2,y2-4,x2+2,y2-4]
        self.d.add(Polygon([v if i%2==0 else self.h-v for i,v in enumerate(pts)],fillColor=HexColor(color),strokeColor=None))
    def panel(self,x,y,label,title):
        self.txt(x,y,label,8,True);self.txt(x+14,y,title,7)
    def axis(self,x,y,w,maximum,ticks,label,top=None):
        for v in ticks:
            xx=x+w*v/maximum
            if top is not None:self.line(xx,top,xx,y,GRID,.35)
            self.line(xx,y,xx,y+3);self.txt(xx,y+13,str(v),6.5,anchor='middle')
        self.line(x,y,x+w,y);self.txt(x+w/2,y+27,label,7,anchor='middle')
    def save(self,name):
        OUT.mkdir(exist_ok=True)
        renderSVG.drawToFile(self.d,str(OUT/f'{name}.svg'))
        svg=OUT/f'{name}.svg'
        content=svg.read_text()
        content=re.sub(r'<svg width="[^"]+" height="[^"]+"',f'<svg width="{W}pt" height="{self.h}pt"',content,count=1)
        content=re.sub(r'viewBox="[^"]+"',f'viewBox="0 0 {W} {self.h}"',content,count=1)
        titles={'fig1_architecture':'学习者状态形成与上下文准备','fig2_temporal':'时序审计及配对迁移','fig3_education':'教育记忆的预算权衡','fig4_locomo':'通用对话证据检索'}
        content=content.replace('<title>...</title>',f'<title>{titles[name]}</title>')
        content=content.replace('<desc>...</desc>','<desc>Original vector figure from retained LearnFlow measurements. See manuscript caption for denominators and statistical scope.</desc>')
        svg.write_text(content)
        from reportlab.pdfgen import canvas
        c=canvas.Canvas(str(OUT/f'{name}.pdf'),pagesize=(W,self.h));c.setTitle(titles[name]);c.setAuthor('')
        renderPDF.draw(self.d,c,0,0);c.showPage();c.save()


def make_architecture_figure():
    f=Figure(330)
    f.panel(8,14,'a','五核按教学问题分工，共用身份、主题与证据来源')
    cores=[('结构 Structure','走到哪里，如何返回','位置 / 依赖 / 锚点'),
           ('知识 Knowledge','理解证据是什么','概念 / 缺口 / 错误'),
           ('人因 Human','当前如何支持','明确请求 / 时效'),
           ('价值 Value','为何学，优先什么','目标 / 确认 / 更新'),
           ('实践 Practice','在何种条件下做成','尝试 / 辅助 / 迁移')]
    for i,(name,q,fields) in enumerate(cores):
        x=8+i*100;f.rect(x,28,92,59,'#F3F6F8',BLUE,.6)
        f.txt(x+46,43,name,7,anchor='middle');f.txt(x+46,60,q,6.5,anchor='middle');f.txt(x+46,75,fields,6.1,anchor='middle')
    f.txt(8,101,'不同核独立判断；学习位置、自述兴趣与受助成功不能互相替代为掌握。',7)
    f.panel(8,122,'b','统一事件更新与版本化记忆')
    labels=[('行为与证据','EvidenceEvent'),('确定性归约','reducer'),('合法状态变更','KernelMutation'),('五核当前状态','KernelState'),('记忆及历史版本','Fact / Module / Claim')]
    for i,(a,b) in enumerate(labels):
        x=8+i*100;f.rect(x,136,92,34,None,'#AAB5BD',.6)
        f.txt(x+46,149,a,7,anchor='middle');f.txt(x+46,161,b,6.3,anchor='middle')
        if i<4:f.arrow(x+93,153,x+99,153)
    f.txt(8,187,'纠正追加证据与后继版本，保留历史；不同核采用各自的内容与巩固门槛。',7)
    f.panel(8,211,'c','按教学用途提供只读证据')
    f.rect(8,223,492,29,'#F3F6F8',BLUE,.6)
    f.txt(254,235,'范围与来源检查 → 当前适用证据 → 预算准入 → ContextPacket',7,anchor='middle')
    f.txt(254,246,'结果、辅助、时间、来源和未知项共同送达',6.5,anchor='middle')
    for x,title,detail in [(8,'Tutor 辅导','相关概念 / 当前支持需求'),(177,'学习规划','目标与路径 / 诊断与独立检查'),(346,'练习与复习','理解与实践证据 / 所属业务对象')]:
        f.arrow(x+77,253,x+77,264);f.rect(x,265,154,31,None,'#AAB5BD',.6)
        f.txt(x+77,278,title,7,anchor='middle');f.txt(x+77,290,detail,6.3,anchor='middle')
    f.txt(8,315,'新的学习行为回到事件入口；讲解、计划和模型生成不直接升级掌握。',7)
    f.save('fig1_architecture')


def make_figures(d):
    make_architecture_figure()

    f=Figure(264)
    f.panel(8,14,'a','全体案例的时序审计')
    order=['formal-02:1800','formal-03:1800','formal-02:3200','formal-03:3200']
    x,w=90,250
    for i,(k,label) in enumerate(zip(order,['1,800 修正前','1,800 修正后','3,200 修正前','3,200 修正后'])):
        y=43+i*34;f.txt(82,y+10,label,7,anchor='end');cum=0
        for status,c in [('pass',BLUE),('fail',STALE),('not_applicable',GRAY)]:
            n=d['freshness'][k].get(status,0)
            if n:
                ww=w*n/1584;f.rect(x+cum,y,ww,18,c,None)
                f.txt(x+cum+ww/2,y+12,str(n),6.5,anchor='middle',color='#FFFFFF' if status in ('pass','fail') else INK);cum+=ww
    f.axis(x,187,w,1584,[0,400,800,1200,1584],'案例数')
    for xx,c,l in [(90,BLUE,'时序一致'),(177,STALE,'过时'),(240,GRAY,'无评估动作 NA')]:
        f.rect(xx,228,7,7,c,None);f.txt(xx+11,234,l,6.5)
    f.panel(366,14,'b','原有过时案例的去向')
    f.txt(431,40,'192 个过时案例',7,anchor='middle')
    # Two explicit bars show paired transition counts; no area illusion.
    for y,n,c,lab in [(73,63,BLUE,'转为时序一致'),(125,129,GRAY,'转为无评估动作')]:
        f.txt(367,y-8,lab,7);f.rect(367,y,131*n/192,19,c,None);f.txt(367+131*n/192+5,y+13,str(n),7)
    f.axis(367,187,131,192,[0,96,192],'配对案例数')
    f.txt(366,240,'仅分析预算 1,800',6.5)
    f.save('fig2_temporal')

    f=Figure(343)
    f.panel(8,14,'a','教育主题及来源 Fact 覆盖')
    arms=['full','facts_only','recent_facts','no_relations','no_guidance','no_episodes','no_bm25','no_aliases','no_fuzzy','no_temporal','no_summary_boost','no_memory']
    x,w=103,204
    for j,arm in enumerate(arms):
        y=46+20*j;f.line(x,y,x+w,y,GRID,.35);f.txt(95,y+2,arm,6.4,anchor='end')
        vals=[100*d['education'][f'{arm}:{b}']['topic_hits']/1794 for b in (1800,3200)]
        f.line(x+w*vals[0]/100,y-1.8,x+w*vals[1]/100,y+1.8,'#ADB6BC',.8)
        f.point(x+w*vals[0]/100,y-1.8,BLUE);f.point(x+w*vals[1]/100,y+1.8,ORANGE,'square')
    f.axis(x,280,w,100,[0,25,50,75,100],'探针送达率（%）')
    f.point(111,327,BLUE);f.txt(118,330,'预算 1,800',6.5)
    f.point(209,327,ORANGE,'square');f.txt(216,330,'预算 3,200',6.5)
    f.panel(337,14,'b','有适用评估案例的动作覆盖')
    f.txt(337,34,'分母 n = 1,311',6.5)
    xx,ww=337,167
    for yy,n,label in [(63,1003,'full  /  1,800'),(122,1311,'full  /  3,200'),(181,0,'no_episodes  /  两档预算')]:
        f.txt(xx,yy-8,label,6.5);f.rect(xx,yy,ww,18,GRAY,None)
        if n:f.rect(xx,yy,ww*n/1311,18,BLUE,None)
        f.txt(xx+ww*n/1311/2 if n else xx+ww/2,yy+12,str(n) if n else '0 / 1,311',6.5,anchor='middle',color='#FFFFFF' if n else INK)
        if 0<n<1311:f.txt(xx+ww*(n+(1311-n)/2)/1311,yy+12,str(1311-n),6.5,anchor='middle')
    f.axis(xx,228,ww,1311,[0,655,1311],'案例数')
    f.rect(337,276,7,7,BLUE,None);f.txt(348,282,'有评估动作',6.5)
    f.rect(422,276,7,7,GRAY,None);f.txt(433,282,'无评估动作',6.5)
    f.txt(337,308,'另有 273 个无适用评估案例',6.5)
    f.txt(337,322,'未进入此分面的分母',6.5)
    f.save('fig3_education')

    f=Figure(263)
    f.panel(8,14,'a','LoCoMo 主类别的证据召回')
    arms=['full','no_bm25','no_aliases','no_fuzzy','no_temporal','recent_facts','no_memory']
    x,w=84,202
    for j,arm in enumerate(arms):
        y=47+22*j;f.line(x,y,x+w,y,GRID,.35);f.txt(76,y+2,arm,6.5,anchor='end')
        vals=[100*d['locomo'][f'{arm}:{b}']['recall'] for b in (1800,3200)]
        f.line(x+w*vals[0]/100,y-1.8,x+w*vals[1]/100,y+1.8,'#ADB6BC',.8)
        f.point(x+w*vals[0]/100,y-1.8,BLUE);f.point(x+w*vals[1]/100,y+1.8,ORANGE,'square')
    f.axis(x,198,w,100,[0,25,50,75,100],'平均证据召回率（%）')
    f.point(91,248,BLUE);f.txt(98,251,'预算 1,800',6.5)
    f.point(194,248,ORANGE,'square');f.txt(201,251,'预算 3,200',6.5)
    f.panel(321,14,'b','BM25 的配对贡献')
    f.txt(321,35,'full − no_bm25；对话等权',6.5)
    xx,ww=337,166
    f.line(xx,53,xx,198,'#87929A',.5,[2,2])
    for yy,b,col,shape in [(85,1800,BLUE,'circle'),(146,3200,ORANGE,'square')]:
        a=next(a for a in d['locomo_paired_intervals'] if a['budget']==b and a['comparison']=='full-minus-no_bm25' and a['metric']=='full_text_recall')
        v=a['cluster_equal_weight_delta']*100;lo,hi=[q*100 for q in a['cluster_bootstrap_95_interval']]
        f.txt(xx+5,yy-24,f'预算 {b:,}',6.5)
        f.line(xx+ww*lo/6,yy,xx+ww*hi/6,yy,col,.8)
        for val in (lo,hi):f.line(xx+ww*val/6,yy-3,xx+ww*val/6,yy+3,col,.8)
        f.point(xx+ww*v/6,yy,col,shape)
        f.txt(xx+ww/2,yy+18,f'{v:.2f}  [{lo:.2f}, {hi:.2f}]',6.5,anchor='middle')
    f.axis(xx,198,ww,6,[0,2,4,6],'配对差（百分点）')
    f.txt(321,250,'10 段对话；2,000 次描述性重采样',6.5)
    f.save('fig4_locomo')


def raster(node):
    script="const fs=require('fs'),path=require('path'),sharp=require('sharp');(async()=>{for(const n of fs.readdirSync(process.argv[1]).filter(x=>x.endsWith('.svg'))){const svg=fs.readFileSync(path.join(process.argv[1],n),'utf8').replace(/(width|height)=\"([0-9.]+)pt\"/g,'$1=\"$2\"');await sharp(Buffer.from(svg),{density:450}).png().toFile(path.join(process.argv[1],n.replace('.svg','.png')));}})().catch(e=>{console.error(e);process.exit(1)});"
    subprocess.run([node,'-e',script,str(OUT)],check=True)


def docx():
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();s=doc.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=Mm(18);s.bottom_margin=Mm(18)
    s.left_margin=s.right_margin=Mm(18.5)
    for name in ('Normal','Title','Heading 1','Heading 2','Caption'):
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(10.5);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Songti SC')
        for e in st.element.findall('.//'+qn('w:rFonts')):
            for k in list(e.attrib):
                if 'theme' in k.lower():del e.attrib[k]
        for e in st.element.findall('.//'+qn('w:pBdr')):e.getparent().remove(e)
    p=doc.styles['Normal'].paragraph_format;p.line_spacing=1.18;p.space_after=Pt(5);p.widow_control=True
    p.first_line_indent=Pt(21)
    doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.first_line_indent=Pt(0);doc.styles['Title'].paragraph_format.space_after=Pt(15)
    for n,size in [('Heading 1',13),('Heading 2',11)]:
        st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True
        st.element.rPr.rFonts.set(qn('w:eastAsia'),'Heiti SC')
        st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(6)
    st=doc.styles['Caption'];st.font.size=Pt(9);st.font.italic=False;st.font.bold=False
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Songti SC')
    st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.12
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.paragraph_format.first_line_indent=Pt(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    lines=(HERE/'manuscript_zh.md').read_text().splitlines();i=0;code=False;ref=False
    while i<len(lines):
        line=lines[i];i+=1
        if not line:continue
        if line.startswith('```'):code=not code;continue
        if code:
            p=doc.add_paragraph(line);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Mm(5)
            p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=not line.startswith('返回')
            for r in p.runs:r.font.size=Pt(9)
        elif line.startswith('# '):
            p=doc.add_paragraph(line[2:].replace('教学的五核','教学的\n五核'),'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_paragraph(line[3:],'Heading 1');ref=line=='## 参考文献'
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('!['):
            file=re.search(r'\]\((.*?)\)',line).group(1)
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE/file),width=Mm(160 if 'fig3_education' in file else 173))
            pic=p._p.xpath('.//wp:docPr')[0];pic.set('descr',line.split(']')[0][2:])
        elif line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
            vals=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.match(r'^\|\s*---',r)]
            n=len(vals[0]);text_table=n==4
            widths=(([26,48,62,37] if vals[0][0]=='状态维度' else [32,45,46,50]) if text_table else ([35,31,38,31,38] if n==5 else ([38,114,21] if vals[0][1]=='干预定义' else [57,58,58])))
            table=doc.add_table(rows=0,cols=n);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
            for col,w in zip(table.columns,widths):col.width=Mm(w)
            # Journal-style three-rule table, per requested scientific visual direction.
            borders=OxmlElement('w:tblBorders')
            for side in ('top','bottom','left','right','insideH','insideV'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if side in ('top','bottom') else 'nil');e.set(qn('w:sz'),'8');e.set(qn('w:color'),'333333');borders.append(e)
            table._tbl.tblPr.append(borders)
            for ri,v in enumerate(vals):
                row=table.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if ri==0:pr.append(OxmlElement('w:tblHeader'))
                for ci,(cell,txt,w) in enumerate(zip(row.cells,v,widths)):
                    cell.width=Mm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(2 if vals[0][1]=='预算 1,800' else 3);p.paragraph_format.space_after=Pt(2 if vals[0][1]=='预算 1,800' else 3);p.paragraph_format.line_spacing=1.04
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or text_table or vals[0][1]=='干预定义' else WD_ALIGN_PARAGRAPH.CENTER
                    r=p.add_run(txt);r.font.size=Pt(9);r.bold=ri==0
                    if ri==0:
                        b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'777777');b.append(e);cell._tc.get_or_add_tcPr().append(b)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(3)
        else:
            nextline=next((v for v in lines[i:] if v.strip()),'')
            cap=bool(re.match(r'^(图|表|算法) [1-9]',line)) and (line.startswith('图 ') or nextline.startswith(('|','```')))
            p=doc.add_paragraph(line,'Caption' if cap else None)
            if cap:
                lead,sep,tail=line.partition('。');p.clear();p.add_run(lead+sep).bold=True
                if tail:p.add_run(tail)
            if cap and line.startswith(('表 ','算法 ')):p.paragraph_format.keep_with_next=True
            if line.startswith('关键词') or ref:p.paragraph_format.first_line_indent=Pt(0)
            if ref:
                p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(6)
                for r in p.runs:r.font.size=Pt(9)
    doc.core_properties.title=lines[0][2:];doc.core_properties.author='';doc.core_properties.subject='五核学习者记忆的教学服务与证据评估'
    path=HERE/'LearnFlow_Educational_Memory_Manuscript_ZH.docx';doc.save(path)
    print('DOCX:',path.name,'tables',len(doc.tables),'figures',len(doc.inline_shapes))


def verify_data(d):
    root=HERE.parents[2]
    for name,digest in d['source_sha256'].items():
        assert hashlib.sha256((root/name).read_bytes()).hexdigest()==digest,name
    audit=json.loads((root/'evals/education_memory_v2/results/formal-03-freshness-v2/summary.json').read_text())
    for g in audit['groups']:
        if g['variant']=='no_episodes':assert g['status_counts']=={'not_applicable':1584}
    assert d['freshness']['formal-03:1800']=={'pass':1003,'not_applicable':581}
    assert d['freshness']['formal-03:3200']=={'pass':1311,'not_applicable':273}
    assert sum(t['cases'] for t in d['freshness_transitions'][0]['transitions'] if t['before']=='fail')==192
    md=(HERE/'manuscript_zh.md').read_text()
    for k,g in d['education'].items():
        arm,b=k.split(':')
        assert g['raw_hits']==0 and g['raw_total']==1152
        a,c=[d['education'][f'{arm}:{v}']['topic_hits'] for v in (1800,3200)]
        assert f'| {arm} | {a:,} / 1,794 | {c:,} / 1,794 |' in md
    for arm in ['full','no_bm25','no_aliases','no_fuzzy','no_temporal','recent_facts','no_memory']:
        a,b=[d['locomo'][f'{arm}:{v}'] for v in (1800,3200)]
        assert f"| {arm} | {a['recall']:.2%} | {a['complete']} / 1,438 | {b['recall']:.2%} | {b['complete']} / 1,438 |" in md
    refs=re.findall(r'^\[(\d+)\]',md,re.M);assert refs==list(map(str,range(1,19)))
    assert '129 个转为无评估动作' in md
    assert '95%' in md
    assert '注册表 2026-09-08.9' in md
    assert '### 2.4 DeepTutor' in md and '3.91/5' in md and '3.80/5' in md
    assert '本文没有运行 TutorBench 或 DeepTutor 基线' in md
    assert re.findall(r'^表 ([1-9]) [^\n]+\n\n\|', md, re.M)==list(map(str,range(1,7)))
    print('Chinese table, negative-result and citation checks passed.')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--docx-only',action='store_true');p.add_argument('--node',default='node');args=p.parse_args()
    d=json.loads((HERE/'derived_results.json').read_text());verify_data(d)
    if not args.docx_only:make_figures(d);raster(args.node)
    docx()
