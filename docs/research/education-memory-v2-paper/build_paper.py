"""Re-derive displays from retained results, then build the editable manuscript.

Run with Python providing python-docx and reportlab; use --data-only to avoid DOCX.
No experiment is executed, no dataset label is modified, and no network is used.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
import gzip
import hashlib
import json
from pathlib import Path
import re

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
RESULTS = ROOT / 'evals/education_memory_v2/results'
INPUTS = {}


def read(relative):
    path = RESULTS / relative
    INPUTS[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix == '.gz':
        with gzip.open(path, 'rt') as handle:
            return [json.loads(line) for line in handle]
    return json.loads(path.read_text())


def derive():
    trials = read('formal-03/education-trials.jsonl.gz')
    assert len(trials) == 38016
    assert len({(r['case_id'], r['budget'], r['variant']) for r in trials}) == 38016
    education = defaultdict(lambda: defaultdict(int))
    for r in trials:
        g = education[f"{r['variant']}:{r['budget']}"]
        g['cases'] += 1
        for name, key in [('topic_hits', 'source_fact_evidence_delivered_formed_source_fact_numerator'),
                          ('topic_total', 'source_fact_evidence_delivered_formed_source_fact_denominator'),
                          ('raw_hits', 'source_fact_evidence_delivered_raw_statement_numerator'),
                          ('raw_total', 'source_fact_evidence_delivered_raw_statement_denominator'),
                          ('episodes', 'episode_count'),
                          ('budget_pass', 'budget')]:
            g[name] += r['metrics'][key] or 0
    expected = {'full': (1407, 1725), 'facts_only': (1416, 1725),
                'recent_facts': (1346, 1656), 'no_relations': (1408, 1725),
                'no_guidance': (1407, 1725), 'no_episodes': (1640, 1793),
                'no_bm25': (1270, 1725), 'no_aliases': (1406, 1725),
                'no_fuzzy': (1413, 1725), 'no_temporal': (1409, 1725),
                'no_summary_boost': (1405, 1725), 'no_memory': (0, 0)}
    manuscript = (HERE / 'manuscript.md').read_text()
    for arm, values in expected.items():
        assert f'| {arm} | {values[0]:,} / 1,794 | {values[1]:,} / 1,794 |' in manuscript
        for budget, hits in zip((1800, 3200), values):
            g = education[f'{arm}:{budget}']
            assert (g['cases'], g['topic_hits'], g['topic_total'], g['raw_hits'], g['raw_total'], g['budget_pass']) == (1584, hits, 1794, 0, 1152, 1584)
    audit = read('formal-03/locomo-primary-audit.json')
    conversational = {}
    for g in audit['current_groups']:
        k = f"{g['variant']}:{g['budget']}"
        assert g['main']['conditions'] == 1438
        conversational[k] = {'recall': g['main']['metrics']['full_text_recall']['mean'],
                             'complete': g['main']['metrics']['full_evidence_complete']['passed'],
                             'multi_hop_complete': g['multi_hop_category_1']['metrics']['full_evidence_complete']['passed'],
                             'zero_recall': g['main']['zero_full_text_recall']}
    for arm in ('full','no_bm25','no_aliases','no_fuzzy','no_temporal','recent_facts','no_memory'):
        a,b = [conversational[f'{arm}:{budget}'] for budget in (1800,3200)]
        assert f"| {arm} | {a['recall']:.2%} | {a['complete']} / 1,438 | {b['recall']:.2%} | {b['complete']} / 1,438 |" in manuscript
    freshness = {}
    for run in ('formal-02','formal-03'):
        d = read(f'{run}-freshness-v2/summary.json')
        assert d['conditions'] == 38016 and d['post_hoc_audit']
        for g in d['groups']:
            if g['variant'] == 'full':
                freshness[f"{run}:{g['budget']}"] = g['status_counts']
    transitions = read('formal-03-freshness-v2/paired-before-after.json')
    assert transitions['groups'][0]['transitions'] == [
        {'before':'fail','after':'not_applicable','cases':129},
        {'before':'fail','after':'pass','cases':63},
        {'before':'not_applicable','after':'not_applicable','cases':452},
        {'before':'pass','after':'pass','cases':940}]
    default = read('default-budget-2900/descriptive-summary.json')
    assert default['cases'] == default['independent_budget_pass'] == 1584
    assert default['actual_practice_action_cases'] == 1311
    final = read('formal-03/final-source-check.json')
    assert final['frozen_source_files'] == 186 and not final['differences']
    result = {'education': dict(education), 'locomo': conversational, 'freshness': freshness,
              'locomo_paired_intervals': audit['main_paired_comparisons'],
              'freshness_transitions': transitions['groups'],
              'default_2900': {k: default[k] for k in ('cases','actual_action_counts','actual_action_freshness','independent_budget_pass')},
              'source_sha256': INPUTS,
              'scope': 'Derived existing measurements only; no new experiment; intervals are retained post hoc descriptive estimates.'}
    (HERE/'derived_results.json').write_text(json.dumps(result, indent=2, ensure_ascii=False)+'\n')
    return result


def figures(d):
    from reportlab.graphics.shapes import Drawing, String, Rect, Line, Polygon
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics import renderSVG
    from reportlab.lib.colors import HexColor
    out = HERE/'figures'; out.mkdir(exist_ok=True)
    blue, orange, gray = [HexColor(c) for c in ('#235789','#C76A2B','#D5DADF')]
    def save(draw, name):
        renderSVG.drawToFile(draw, str(out/f'{name}.svg'))
    arch=Drawing(680,270)
    arch.add(String(15,252,'AUTHORITATIVE FORMATION',fontName='Helvetica-Bold',fontSize=12))
    labels=['Evidence\nEvent','Deterministic\nreducer','Kernel\nMutation','Kernel\nState','Fact > Module\n> Claim']
    for i,label in enumerate(labels):
        x=15+135*i
        arch.add(Rect(x,180,115,53,fillColor=HexColor('#F1F4F7'),strokeColor=blue))
        for j,s in enumerate(label.split('\n')):
            arch.add(String(x+57.5,212-j*15,s,textAnchor='middle',fontName='Helvetica',fontSize=11))
        if i<4:
            arch.add(Line(x+116,206,x+130,206,strokeColor=blue))
            arch.add(Polygon([x+130,206,x+125,209,x+125,203],fillColor=blue,strokeColor=blue))
    arch.add(String(15,143,'READ ONLY PREPARATION',fontName='Helvetica-Bold',fontSize=12))
    for x,w,txt in [(15,190,'Scope and provenance checks'),(242,190,'Budgeted episode prefix'),(470,195,'Deterministic plan actions')]:
        arch.add(Rect(x,74,w,44,fillColor=HexColor('#FFFFFF'),strokeColor=blue))
        arch.add(String(x+w/2,91,txt,textAnchor='middle',fontSize=11))
    for x in (207,434):
        arch.add(Line(x,96,x+30,96,strokeColor=blue))
        arch.add(Polygon([x+30,96,x+25,99,x+25,93],fillColor=blue,strokeColor=blue))
    arch.add(String(15,40,'Validated attempts and live controls enter as scoped evidence.',fontSize=11))
    arch.add(String(15,21,'Plans and summaries do not independently establish mastery.',fontSize=11))
    save(arch,'architecture')
    fresh=Drawing(660,350)
    chart=VerticalBarChart();chart.x=57;chart.y=67;chart.width=570;chart.height=232
    order=['formal-02:1800','formal-03:1800','formal-02:3200','formal-03:3200']
    chart.data=[[d['freshness'][k].get(s,0) for k in order] for s in ('pass','fail','not_applicable')]
    chart.categoryAxis.categoryNames=['Before 1800','After 1800','Before 3200','After 3200']
    chart.categoryAxis.style='stacked';chart.valueAxis.valueMin=0;chart.valueAxis.valueMax=1600;chart.valueAxis.valueStep=400
    chart.bars[0].fillColor=blue;chart.bars[1].fillColor=orange;chart.bars[2].fillColor=gray
    fresh.add(chart)
    for i,key in enumerate(order):
        cumulative=0
        for status in ('pass','fail','not_applicable'):
            count=d['freshness'][key].get(status,0)
            if count:
                fresh.add(String(chart.x+chart.width*(i+.5)/4,
                                 chart.y+chart.height*(cumulative+count/2)/1600-4,
                                 str(count),textAnchor='middle',fontName='Helvetica-Bold',fontSize=11,
                                 fillColor=HexColor('#FFFFFF' if status=='pass' else '#111111')))
            cumulative+=count
    fresh.add(String(57,325,'Temporal audit outcome counts',fontName='Helvetica-Bold',fontSize=13))
    for x,c,label in [(57,blue,'Consistent action'),(252,orange,'Stale action'),(417,gray,'No assessment action')]:
        fresh.add(Rect(x,21,12,12,fillColor=c,strokeColor=c));fresh.add(String(x+19,23,label,fontSize=10))
    save(fresh,'freshness')
    coverage=Drawing(690,330)
    for x,width,arms,source,title in [(47,275,['full','facts_only','no_episodes'],'education','Educational topic probes'),(395,258,['full','no_bm25'],'locomo','LoCoMo evidence recall')]:
        c=VerticalBarChart();c.x=x;c.y=65;c.width=width;c.height=205
        c.data=[[100*(d[source][f'{arm}:{b}']['topic_hits']/1794 if source=='education' else d[source][f'{arm}:{b}']['recall']) for arm in arms] for b in (1800,3200)]
        c.categoryAxis.categoryNames=[a.replace('_',' ') for a in arms]
        c.valueAxis.valueMin=0;c.valueAxis.valueMax=100;c.valueAxis.valueStep=25
        c.bars[0].fillColor=blue;c.bars[1].fillColor=orange
        c.barLabelFormat='%.1f';c.barLabels.fontSize=9;c.barLabels.dy=7
        coverage.add(c);coverage.add(String(x,293,title+' (%)',fontName='Helvetica-Bold',fontSize=12))
    for x,c,label in [(195,blue,'Budget 1800'),(365,orange,'Budget 3200')]:
        coverage.add(Rect(x,17,12,12,fillColor=c,strokeColor=c));coverage.add(String(x+19,19,label,fontSize=11))
    save(coverage,'coverage')


def build_docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.75)
    sec.left_margin=sec.right_margin=Inches(.85)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(11)
        for fonts in st.element.findall('.//' + qn('w:rFonts')):
            for key in list(fonts.attrib):
                if 'theme' in key.lower(): del fonts.attrib[key]
        for border in st.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=1.10;normal.space_after=Pt(6);normal.widow_control=True
    doc.styles['Title'].font.size=Pt(20)
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    for name,size in [('Heading 1',14),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(12)
        doc.styles[name].paragraph_format.space_after=Pt(6)
    doc.styles['Caption'].font.size=Pt(10)
    doc.styles['Caption'].font.bold=False
    h=sec.header.paragraphs[0];h.text='Educational agent memory   |   Research manuscript draft';h.style='Caption'
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    lines=(HERE/'manuscript.md').read_text().splitlines();i=0;code=False
    # Keep table captions above their table, including captions written below
    # the Markdown table for source readability.
    moved_captions=set()
    while i<len(lines):
        if i in moved_captions:
            i+=1;continue
        line=lines[i];i+=1
        if not line.strip():continue
        if line.startswith('```'):
            code=not code;continue
        if code:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.keep_with_next=not line.startswith('Return')
            r=p.add_run(line);r.font.name='Courier New';r.font.size=Pt(9)
        elif line.startswith('!['):
            path=re.search(r'\]\((.*?)\)',line).group(1)
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(HERE/path),width=Inches(6.65))
        elif line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
            look=i
            while look<len(lines) and not lines[look].strip():look+=1
            if look<len(lines) and re.match(r'^Table \d+\.',lines[look]):
                p=doc.add_paragraph(lines[look],'Caption')
                p.paragraph_format.keep_with_next=True
                moved_captions.add(look)
            cells=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.match(r'^\|\s*---',r)]
            table=doc.add_table(rows=0,cols=len(cells[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.autofit=False
            widths={2:[3.35,3.35],3:[2.05,2.325,2.325],5:[1.65,1.05,1.45,1.05,1.5]}[len(cells[0])]
            if cells[0][0]=='Configuration' and cells[0][-1]=='Track':widths=[1.6,4.15,.95]
            for col,w in zip(table.columns,widths):col.width=Inches(w)
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            table._tbl.tblPr.append(borders)
            for ri,values in enumerate(cells):
                row=table.add_row()
                trpr=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
                if ri==0:trpr.append(OxmlElement('w:tblHeader'))
                for ci,(cell,txt,w) in enumerate(zip(row.cells,values,widths)):
                    cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.line_spacing=1.05
                    if ci>0 and len(cells[0])==5:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    run=p.add_run(txt);run.font.size=Pt(10);run.bold=ri==0
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF2');cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        else:
            caption=bool(re.match(r'^(Figure|Table) \d+\.',line))
            p=doc.add_paragraph(line,'Caption' if caption else None)
            if caption and i<len(lines) and lines[i:i+2] and any(l.startswith('|') for l in lines[i:i+2]):p.paragraph_format.keep_with_next=True
            if re.match(r'^\[\d+\]',line):
                p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(10)
    doc.core_properties.title='Temporal Validity and Evidence Coverage in Budgeted Memory for Educational Agents'
    doc.core_properties.subject='Evidence-grounded systems study; author information not supplied'
    doc.core_properties.author=''
    doc.save(HERE/'LearnFlow_Educational_Memory_Manuscript.docx')
    print('DOCX paragraphs:',len(doc.paragraphs),'tables:',len(doc.tables))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--data-only',action='store_true');parser.add_argument('--docx-only',action='store_true');args=parser.parse_args()
    if not args.docx_only:
        data=derive();figures(data);print('All table assertions passed; source-derived JSON and SVG figures written.')
    if not args.data_only:build_docx()
